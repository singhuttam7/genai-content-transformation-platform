from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.ingestion.parsers.docx import DOCXProcessor
from app.ingestion.schemas import (
    ContentBlockType,
    IngestionRequest,
    InputType,
)


class InMemoryResolver:
    """Test resolver that returns predefined binary content."""

    def __init__(self, content: bytes) -> None:
        self.content = content

    async def resolve(self, request: IngestionRequest) -> bytes:
        return self.content


def create_docx(
    *,
    include_heading: bool = True,
    include_table: bool = True,
    include_list: bool = True,
) -> bytes:
    """Create a DOCX fixture entirely in memory."""

    document = Document()

    if include_heading:
        document.add_heading(
            "Cybersecurity Advisory",
            level=1,
        )

    document.add_paragraph(
        "This is a sample cybersecurity advisory "
        "for processor testing."
    )

    if include_list:
        document.add_paragraph(
            "First security recommendation",
            style="List Bullet",
        )
        document.add_paragraph(
            "Second security recommendation",
            style="List Bullet",
        )

    if include_table:
        table = document.add_table(
            rows=2,
            cols=2,
        )

        table.cell(0, 0).text = "Threat"
        table.cell(0, 1).text = "Severity"

        table.cell(1, 0).text = "Credential Theft"
        table.cell(1, 1).text = "High"

    document.add_paragraph(
        "Organizations should review the affected systems."
    )

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def create_empty_docx() -> bytes:
    """Create a DOCX containing no meaningful text."""

    document = Document()

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


@pytest.mark.asyncio
async def test_docx_processor_success() -> None:
    content = create_docx()

    request = IngestionRequest(
        input_type=InputType.DOCX,
        title="Security Advisory",
        filename="advisory.docx",
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        content=content,
    )

    resolver = InMemoryResolver(content)

    processor = DOCXProcessor(
        resolver=resolver,
    )

    result = await processor.process(request)

    assert result.source.source_type == InputType.DOCX
    assert result.title == "Security Advisory"

    assert result.text.strip()
    assert len(result.blocks) > 0

    assert result.metadata["processor"] == "docx"
    assert result.metadata["block_count"] == len(result.blocks)


@pytest.mark.asyncio
async def test_docx_heading_extraction() -> None:
    content = create_docx(
        include_heading=True,
        include_table=False,
        include_list=False,
    )

    request = IngestionRequest(
        input_type=InputType.DOCX,
        title="Heading Test",
        filename="heading.docx",
        content=content,
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(content),
    )

    result = await processor.process(request)

    headings = [
        block
        for block in result.blocks
        if block.block_type == ContentBlockType.HEADING
    ]

    assert len(headings) == 1
    assert headings[0].content == "Cybersecurity Advisory"
    assert headings[0].metadata["heading_level"] == 1


@pytest.mark.asyncio
async def test_docx_list_extraction() -> None:
    content = create_docx(
        include_heading=False,
        include_table=False,
        include_list=True,
    )

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="list.docx",
        content=content,
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(content),
    )

    result = await processor.process(request)

    list_blocks = [
        block
        for block in result.blocks
        if block.block_type == ContentBlockType.LIST
    ]

    assert len(list_blocks) == 2

    assert list_blocks[0].content == (
        "First security recommendation"
    )

    assert list_blocks[1].content == (
        "Second security recommendation"
    )


@pytest.mark.asyncio
async def test_docx_table_extraction() -> None:
    content = create_docx(
        include_heading=False,
        include_table=True,
        include_list=False,
    )

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="table.docx",
        content=content,
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(content),
    )

    result = await processor.process(request)

    tables = [
        block
        for block in result.blocks
        if block.block_type == ContentBlockType.TABLE
    ]

    assert len(tables) == 1

    table = tables[0]

    assert table.metadata["rows"] == 2
    assert table.metadata["columns"] == 2

    assert "Threat" in table.content
    assert "Severity" in table.content
    assert "Credential Theft" in table.content
    assert "High" in table.content


@pytest.mark.asyncio
async def test_docx_preserves_element_order() -> None:
    content = create_docx(
        include_heading=True,
        include_table=True,
        include_list=False,
    )

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="order.docx",
        content=content,
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(content),
    )

    result = await processor.process(request)

    orders = [
        block.order
        for block in result.blocks
    ]

    assert orders == list(range(len(result.blocks)))

    block_types = [
        block.block_type
        for block in result.blocks
    ]

    assert block_types[0] == ContentBlockType.HEADING
    assert block_types[1] == ContentBlockType.PARAGRAPH
    assert block_types[2] == ContentBlockType.TABLE


@pytest.mark.asyncio
async def test_docx_empty_content() -> None:
    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="empty.docx",
        content=b"",
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(b""),
    )

    with pytest.raises(
        ValueError,
        match="DOCX content cannot be empty",
    ):
        await processor.process(request)


@pytest.mark.asyncio
async def test_docx_corrupted_content() -> None:
    corrupted_content = b"This is not a valid DOCX file."

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="corrupted.docx",
        content=corrupted_content,
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(corrupted_content),
    )

    with pytest.raises(
        ValueError,
        match="Unable to read DOCX document",
    ):
        await processor.process(request)


@pytest.mark.asyncio
async def test_docx_empty_document() -> None:
    content = create_empty_docx()

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="empty-document.docx",
        content=content,
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(content),
    )

    with pytest.raises(
        ValueError,
        match="no extractable textual content",
    ):
        await processor.process(request)


@pytest.mark.asyncio
async def test_docx_wrong_input_type() -> None:
    content = create_docx()

    request = IngestionRequest(
        input_type=InputType.PDF,
        filename="wrong-type.docx",
        content=content,
    )

    processor = DOCXProcessor(
        resolver=InMemoryResolver(content),
    )

    with pytest.raises(
        ValueError,
        match="can only process DOCX inputs",
    ):
        await processor.process(request)


@pytest.mark.asyncio
async def test_docx_resolver_failure() -> None:
    class FailingResolver:
        async def resolve(
            self,
            request: IngestionRequest,
        ) -> bytes:
            raise RuntimeError("resolver failure")

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="resolver-failure.docx",
        content=b"placeholder",
    )

    processor = DOCXProcessor(
        resolver=FailingResolver(),
    )

    with pytest.raises(
        RuntimeError,
        match="resolver failure",
    ):
        await processor.process(request)