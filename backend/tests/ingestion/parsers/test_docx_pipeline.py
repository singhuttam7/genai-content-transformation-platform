from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.ingestion.detectors import DefaultInputDetector
from app.ingestion.normalization import DefaultContentNormalizer
from app.ingestion.pipeline import IngestionPipeline
from app.ingestion.registry import create_processor_router
from app.ingestion.schemas import (
    ContentBlockType,
    IngestionRequest,
    InputType,
)


def create_docx_fixture() -> bytes:
    """Create a representative DOCX document in memory."""

    document = Document()

    document.add_heading(
        "Cybersecurity Threat Advisory",
        level=1,
    )

    document.add_paragraph(
        "A critical vulnerability has been identified "
        "in a network appliance."
    )

    document.add_heading(
        "Recommended Actions",
        level=2,
    )

    document.add_paragraph(
        "Apply the latest security patch.",
        style="List Bullet",
    )

    document.add_paragraph(
        "Review authentication logs.",
        style="List Bullet",
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Risk"
    table.cell(0, 1).text = "Severity"

    table.cell(1, 0).text = "Remote Exploitation"
    table.cell(1, 1).text = "Critical"

    document.add_paragraph(
        "Organizations should prioritize remediation "
        "of affected systems."
    )

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


@pytest.mark.asyncio
async def test_docx_complete_ingestion_pipeline() -> None:
    """Test DOCX from detection through normalization."""

    content = create_docx_fixture()

    request = IngestionRequest(
        input_type=InputType.DOCX,
        title="Cybersecurity Threat Advisory",
        filename="threat-advisory.docx",
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        content=content,
    )

    detector = DefaultInputDetector()
    router = create_processor_router()
    normalizer = DefaultContentNormalizer()

    pipeline = IngestionPipeline(
        detector=detector,
        router=router,
        normalizer=normalizer,
    )

    result = await pipeline.run(request)

    assert result.source.source_type == InputType.DOCX

    assert result.title == (
        "Cybersecurity Threat Advisory"
    )

    assert result.text.strip()

    assert len(result.segments) > 0

    assert result.provenance["ingestion_normalizer"] == (
        "default"
    )

    assert result.provenance["normalization_version"] == (
        "1.0"
    )


@pytest.mark.asyncio
async def test_docx_structure_survives_pipeline() -> None:
    """Verify important DOCX structures survive processing."""

    content = create_docx_fixture()

    request = IngestionRequest(
        input_type=InputType.DOCX,
        title="Structure Test",
        filename="structure.docx",
        content=content,
    )

    detector = DefaultInputDetector()
    router = create_processor_router()
    normalizer = DefaultContentNormalizer()

    pipeline = IngestionPipeline(
        detector=detector,
        router=router,
        normalizer=normalizer,
    )

    result = await pipeline.run(request)

    assert "Cybersecurity Threat Advisory" in result.text
    assert "Recommended Actions" in result.text
    assert "Apply the latest security patch." in result.text
    assert "Review authentication logs." in result.text
    assert "Remote Exploitation" in result.text
    assert "Critical" in result.text

    assert len(result.segments) >= 5


@pytest.mark.asyncio
async def test_docx_detector_router_processor_integration() -> None:
    """
    Verify that DOCX is correctly detected and routed to
    DOCXProcessor before reaching the normalizer.
    """

    content = create_docx_fixture()

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="integration.docx",
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        content=content,
    )

    detector = DefaultInputDetector()

    detected_type = detector.detect(request)

    assert detected_type == InputType.DOCX

    router = create_processor_router()

    assert router.supports(InputType.DOCX)

    processor = router.get_processor(InputType.DOCX)

    assert type(processor).__name__ == "DOCXProcessor"

    normalizer = DefaultContentNormalizer()

    pipeline = IngestionPipeline(
        detector=detector,
        router=router,
        normalizer=normalizer,
    )

    result = await pipeline.run(request)

    assert result.source.source_type == InputType.DOCX

    assert result.text.strip()


@pytest.mark.asyncio
async def test_docx_pipeline_normalizes_whitespace() -> None:
    """Verify DOCX content reaches the normalizer correctly."""

    document = Document()

    document.add_heading(
        "Whitespace Test",
        level=1,
    )

    document.add_paragraph(
        "First paragraph."
    )

    document.add_paragraph(
        "Second paragraph."
    )

    buffer = BytesIO()
    document.save(buffer)

    content = buffer.getvalue()

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="whitespace.docx",
        content=content,
    )

    pipeline = IngestionPipeline(
        detector=DefaultInputDetector(),
        router=create_processor_router(),
        normalizer=DefaultContentNormalizer(),
    )

    result = await pipeline.run(request)

    assert result.text.strip()

    assert "First paragraph." in result.text
    assert "Second paragraph." in result.text

    # Canonical segments are structured ContentBlock objects.
    assert all(
        segment.content.strip()
        for segment in result.segments
    )


@pytest.mark.asyncio
async def test_docx_pipeline_preserves_content_categories() -> None:
    """Verify headings, lists, and tables are extracted."""

    content = create_docx_fixture()

    request = IngestionRequest(
        input_type=InputType.DOCX,
        filename="categories.docx",
        content=content,
    )

    detector = DefaultInputDetector()
    router = create_processor_router()
    normalizer = DefaultContentNormalizer()

    processor = router.get_processor(
        InputType.DOCX
    )

    extracted = await processor.process(request)

    block_types = {
        block.block_type
        for block in extracted.blocks
    }

    assert ContentBlockType.HEADING in block_types
    assert ContentBlockType.PARAGRAPH in block_types
    assert ContentBlockType.LIST in block_types
    assert ContentBlockType.TABLE in block_types

    pipeline = IngestionPipeline(
        detector=detector,
        router=router,
        normalizer=normalizer,
    )

    canonical = await pipeline.run(request)

    assert canonical.text.strip()
    assert len(canonical.segments) > 0

    canonical_block_types = {
        block.block_type
        for block in canonical.segments
    }

    assert ContentBlockType.HEADING in canonical_block_types
    assert ContentBlockType.PARAGRAPH in canonical_block_types
    assert ContentBlockType.LIST in canonical_block_types
    assert ContentBlockType.TABLE in canonical_block_types